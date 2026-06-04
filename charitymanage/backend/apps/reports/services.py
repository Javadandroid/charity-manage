import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import jdatetime
from django.db.models import Sum, F

from apps.invoices.models import Invoice, InvoiceItem
from apps.booths.models import Booth
from apps.products.models import Product

class ReportService:
    """لایه سرویس برای تولید گزارشات اکسل"""
    
    @classmethod
    def _get_styles(cls):
        return {
            'header_font': Font(name='Estedad Regular', size=12, bold=True),
            'normal_font': Font(name='Estedad Regular', size=11),
            'header_fill': PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid'),
            'thin_border': Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            ),
            'center_alignment': Alignment(horizontal='center', vertical='center'),
            'right_alignment': Alignment(horizontal='right', vertical='center'),
            'number_format_thousands': '#,##0'
        }

    @classmethod
    def generate_market_report(cls) -> bytes:
        wb = openpyxl.Workbook()
        styles = cls._get_styles()

        # --- شیت اول: فروش کل ---
        ws1 = wb.active
        ws1.title = "فروش کل"
        ws1.sheet_view.rightToLeft = True

        column_widths_ws1 = {'A': 5, 'B': 18, 'C': 25, 'D': 15, 'E': 10, 'F': 18, 'G': 15, 'H': 20}
        for col_letter, width in column_widths_ws1.items():
            ws1.column_dimensions[col_letter].width = width

        headers_ws1 = ['ردیف', 'تاریخ', 'نام کالا', 'قیمت واحد', 'تعداد', 'قیمت کل', 'شماره فاکتور', 'غرفه']
        for col_num, header_title in enumerate(headers_ws1, 1):
            cell = ws1.cell(row=1, column=col_num, value=header_title)
            cell.font = styles['header_font']
            cell.fill = styles['header_fill']
            cell.border = styles['thin_border']
            cell.alignment = styles['center_alignment']

        invoice_items = InvoiceItem.objects.select_related('invoice', 'product', 'invoice__booth').order_by('invoice__created_at')

        for row_idx, item in enumerate(invoice_items, 1):
            jalali_datetime = jdatetime.datetime.fromgregorian(datetime=item.invoice.created_at)
            jalali_date_str = jalali_datetime.strftime('%Y/%m/%d %H:%M')

            row_data = [
                row_idx,
                jalali_date_str,
                item.product.name,
                item.price,
                item.quantity,
                item.price * item.quantity,
                item.invoice.invoice_number,
                item.invoice.booth.name,
            ]

            for col_idx, cell_value in enumerate(row_data, 1):
                cell = ws1.cell(row=row_idx + 1, column=col_idx, value=cell_value)
                cell.font = styles['normal_font']
                cell.border = styles['thin_border']
                cell.alignment = styles['center_alignment']
                if col_idx in [4, 6]:
                    cell.number_format = styles['number_format_thousands']

        # --- شیت دوم: فروش روزانه ---
        ws2 = wb.create_sheet(title="فروش روزانه")
        ws2.sheet_view.rightToLeft = True

        all_invoices = Invoice.objects.all().order_by('created_at')
        jalali_dates_for_header = []
        for inv in all_invoices:
            j_date = jdatetime.date.fromgregorian(date=inv.created_at.date())
            jalali_dates_for_header.append(j_date.strftime('%Y/%m/%d'))

        unique_jalali_dates = sorted(list(set(jalali_dates_for_header)))
        all_booths = Booth.objects.all().order_by('name')

        ws2.column_dimensions['A'].width = 25
        for i, _ in enumerate(unique_jalali_dates):
            col_letter = get_column_letter(i + 2)
            ws2.column_dimensions[col_letter].width = 15
        ws2.column_dimensions[get_column_letter(len(unique_jalali_dates) + 2)].width = 18

        headers_ws2 = ['نام غرفه'] + unique_jalali_dates + ['فروش کل']
        for col_num, header_title in enumerate(headers_ws2, 1):
            cell = ws2.cell(row=1, column=col_num, value=header_title)
            cell.font = styles['header_font']
            cell.fill = styles['header_fill']
            cell.border = styles['thin_border']
            cell.alignment = styles['center_alignment']

        current_row_num_ws2 = 2
        grand_total_sales = 0
        daily_totals_for_footer = {date_str: 0 for date_str in unique_jalali_dates}

        for booth_obj in all_booths:
            cell_booth_name = ws2.cell(row=current_row_num_ws2, column=1, value=booth_obj.name)
            cell_booth_name.font = styles['normal_font']
            cell_booth_name.border = styles['thin_border']
            cell_booth_name.alignment = styles['center_alignment']

            total_booth_sales_for_row = 0
            for col_idx, date_str in enumerate(unique_jalali_dates, 2):
                date_parts = date_str.split('/')
                j_date_obj = jdatetime.date(int(date_parts[0]), int(date_parts[1]), int(date_parts[2]))
                g_date_obj = j_date_obj.togregorian()

                daily_sales_for_booth = Invoice.objects.filter(
                    booth_id=booth_obj.id,
                    created_at__date=g_date_obj
                ).aggregate(total_sales=Sum(F('total_amount')))['total_sales'] or 0

                cell = ws2.cell(row=current_row_num_ws2, column=col_idx, value=daily_sales_for_booth)
                cell.font = styles['normal_font']
                cell.border = styles['thin_border']
                cell.alignment = styles['center_alignment']
                cell.number_format = styles['number_format_thousands']

                total_booth_sales_for_row += daily_sales_for_booth
                daily_totals_for_footer[date_str] += daily_sales_for_booth

            cell_total_booth_sales = ws2.cell(row=current_row_num_ws2, column=len(unique_jalali_dates) + 2, value=total_booth_sales_for_row)
            cell_total_booth_sales.font = Font(name='Estedad Regular', size=11, bold=True)
            cell_total_booth_sales.border = styles['thin_border']
            cell_total_booth_sales.alignment = styles['center_alignment']
            cell_total_booth_sales.fill = PatternFill(start_color='FFEBF5', end_color='FFEBF5', fill_type='solid')
            cell_total_booth_sales.number_format = styles['number_format_thousands']
            grand_total_sales += total_booth_sales_for_row
            current_row_num_ws2 += 1

        footer_row_num = current_row_num_ws2
        cell_footer_label = ws2.cell(row=footer_row_num, column=1, value="جمع کل فروش")
        cell_footer_label.font = styles['header_font']
        cell_footer_label.border = styles['thin_border']
        cell_footer_label.alignment = styles['center_alignment']
        cell_footer_label.fill = PatternFill(start_color='FFEBF5', end_color='FFEBF5', fill_type='solid')

        for col_idx, date_str in enumerate(unique_jalali_dates, 2):
            cell_daily_total = ws2.cell(row=footer_row_num, column=col_idx, value=daily_totals_for_footer[date_str])
            cell_daily_total.font = Font(name='Estedad Regular', size=11, bold=True)
            cell_daily_total.border = styles['thin_border']
            cell_daily_total.alignment = styles['center_alignment']
            cell_daily_total.fill = PatternFill(start_color='FFEBF5', end_color='FFEBF5', fill_type='solid')
            cell_daily_total.number_format = styles['number_format_thousands']

        cell_grand_total = ws2.cell(row=footer_row_num, column=len(unique_jalali_dates) + 2, value=grand_total_sales)
        cell_grand_total.font = Font(name='Estedad Regular', size=12, bold=True)
        cell_grand_total.border = styles['thin_border']
        cell_grand_total.alignment = styles['center_alignment']
        cell_grand_total.fill = PatternFill(start_color='FFEBF5', end_color='FFEBF5', fill_type='solid')
        cell_grand_total.number_format = styles['number_format_thousands']

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @classmethod
    def generate_product_sales_report(cls) -> bytes:
        wb = openpyxl.Workbook()
        styles = cls._get_styles()
        ws = wb.active
        ws.title = "فروش محصولات"
        ws.sheet_view.rightToLeft = True

        column_widths = {'A': 5, 'B': 40, 'C': 15, 'D': 18, 'E': 25}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        headers = ['ردیف', 'نام محصول', 'تعداد فروش', 'قیمت واحد', 'غرفه']
        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header_title)
            cell.font = styles['header_font']
            cell.fill = styles['header_fill']
            cell.border = styles['thin_border']
            cell.alignment = styles['center_alignment']

        products = Product.objects.select_related('booth', 'unit').all().order_by('booth__name', 'name')

        row_idx = 1
        for product_item in products:
            row_idx += 1
            total_sold = product_item.invoice_items.aggregate(sold_quantity=Sum('quantity'))['sold_quantity'] or 0
            unit_name = f" ({product_item.unit.name})" if product_item.unit else ""
            product_name_with_unit = f"{product_item.name}{unit_name}"

            row_data = [
                row_idx - 1,
                product_name_with_unit,
                total_sold,
                product_item.price,
                product_item.booth.name if product_item.booth else "نامشخص",
            ]

            for col_idx, cell_value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
                cell.font = styles['normal_font']
                cell.border = styles['thin_border']
                if col_idx in [1, 3, 4]:
                    cell.alignment = styles['center_alignment']
                else:
                    cell.alignment = styles['right_alignment']
                
                if col_idx == 4:
                    cell.number_format = styles['number_format_thousands']
                elif col_idx == 3:
                    cell.number_format = '0'

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def _normalize_phone(phone_str):
        if not phone_str:
            return ""
        persian_to_english = {
            '۰': '0', '۱': '1', '۲': '2', '۳': '3', '۴': '4',
            '۵': '5', '۶': '6', '۷': '7', '۸': '8', '۹': '9',
            '٠': '0', '١': '1', '٢': '2', '٣': '3', '٤': '4',
            '٥': '5', '٦': '6', '٧': '7', '٨': '8', '٩': '9',
        }
        cleaned = "".join(persian_to_english.get(char, char) for char in str(phone_str))
        cleaned = "".join(c for c in cleaned if c.isdigit())
        if len(cleaned) == 10 and cleaned.startswith('9'):
            cleaned = '0' + cleaned
        return cleaned

    @staticmethod
    def _consolidate_names(names_set):
        names = list(names_set)
        cleaned_names = [n for n in names if n not in ["بدون نام", "کیوسک", ""]]
        if cleaned_names:
            return " / ".join(cleaned_names)
        elif names:
            non_empty = [n for n in names if n]
            if non_empty:
                return " / ".join(non_empty)
        return "بدون نام"

    @classmethod
    def generate_customer_phones_report(cls) -> bytes:
        invoices = Invoice.objects.exclude(customer_phone='').exclude(customer_phone__isnull=True)
        customers = {}
        for inv in invoices:
            phone = cls._normalize_phone(inv.customer_phone)
            if not phone:
                continue
            
            if phone not in customers:
                customers[phone] = {
                    'names': set(),
                    'invoice_numbers': [],
                    'total_amount': 0
                }
            
            name = inv.customer_name.strip() if inv.customer_name else ""
            if name:
                customers[phone]['names'].add(name)
            
            inv_num_raw = inv.invoice_number
            inv_num_clean = "".join(c for c in inv_num_raw if c.isdigit())
            inv_num_display = str(int(inv_num_clean)) if inv_num_clean else inv_num_raw
                
            customers[phone]['invoice_numbers'].append(inv_num_display)
            customers[phone]['total_amount'] += inv.total_amount
            
        rows = []
        for phone, data in customers.items():
            name = cls._consolidate_names(data['names'])
            try:
                sorted_inv_nums = sorted(data['invoice_numbers'], key=int)
            except ValueError:
                sorted_inv_nums = sorted(data['invoice_numbers'])
            invoice_nums_str = ", ".join(sorted_inv_nums)
            
            rows.append({
                'phone': phone,
                'name': name,
                'invoice_numbers': invoice_nums_str,
                'total_amount': data['total_amount']
            })
            
        rows.sort(key=lambda x: x['total_amount'], reverse=True)
        
        wb = openpyxl.Workbook()
        styles = cls._get_styles()
        ws = wb.active
        ws.title = "شماره تماس مشتریان"
        ws.sheet_view.rightToLeft = True
        
        column_widths = {'A': 8, 'B': 25, 'C': 18, 'D': 35, 'E': 20}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width
            
        headers = ['ردیف', 'نام', 'شماره', 'شماره‌های فاکتور', 'مبلغ کل خرید']
        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header_title)
            cell.font = styles['header_font']
            cell.fill = styles['header_fill']
            cell.border = styles['thin_border']
            cell.alignment = styles['center_alignment']
            
        for idx, row_data in enumerate(rows, 1):
            row_idx = idx + 1
            
            ws.cell(row=row_idx, column=1, value=idx)
            ws.cell(row=row_idx, column=2, value=row_data['name'])
            ws.cell(row=row_idx, column=3, value=row_data['phone'])
            ws.cell(row=row_idx, column=4, value=row_data['invoice_numbers'])
            ws.cell(row=row_idx, column=5, value=row_data['total_amount'])
            
            for col_idx in range(1, 6):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = styles['normal_font']
                cell.border = styles['thin_border']
                if col_idx in [1, 3, 4, 5]:
                    cell.alignment = styles['center_alignment']
                else:
                    cell.alignment = styles['right_alignment']
                    
                if col_idx == 5:
                    cell.number_format = styles['number_format_thousands']
                elif col_idx == 3:
                    cell.number_format = '@'
                    
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
