from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils.translation import gettext_lazy as _  # global gettext_lazy
from django.db.models import Sum
from django.http import JsonResponse, HttpResponse
import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import Product, UnitOfMeasure
from .forms import ProductForm, UnitOfMeasureForm, ProductBulkUploadForm
from booths.models import Booth
from django.db.models import F
# Create your views here.

@login_required
def product_list(request):
    """نمایش لیست محصولات"""
    # اگر کاربر ادمین باشد، تمام محصولات را نشان دهید
    if request.user.is_staff:
        products = Product.objects.all()
        all_booths = Booth.objects.all()
    else:
        # در غیر این صورت، فقط محصولات غرفه‌هایی که کاربر مدیر یا کارمند آنهاست را نشان دهید
        managed_booths = Booth.objects.filter(manager=request.user)
        staff_booths = Booth.objects.filter(staff=request.user)
        booths = (managed_booths | staff_booths).distinct()
        products = Product.objects.filter(booth__in=booths)
        all_booths = booths
    
    # اعمال فیلترهای جستجو
    if 'name' in request.GET and request.GET['name']:
        products = products.filter(name__icontains=request.GET['name'])
    
    if 'code' in request.GET and request.GET['code']:
        products = products.filter(code__icontains=request.GET['code'])
    
    if 'booth' in request.GET and request.GET['booth']:
        products = products.filter(booth_id=request.GET['booth'])
    
    if 'is_available' in request.GET and request.GET['is_available'] != '':
        is_available = request.GET['is_available'] == '1'
        products = products.filter(is_available=is_available)
    
    context = {
        'products': products,
        'title': _('محصولات'),
        'all_booths': all_booths
    }
    return render(request, 'products/product_list.html', context)

@login_required
def product_detail(request, pk):
    """نمایش جزئیات یک محصول"""
    product = get_object_or_404(Product, pk=pk)
    
    # بررسی دسترسی کاربر به محصول
    if not request.user.is_staff and request.user != product.booth.manager and request.user not in product.booth.staff.all():
        messages.error(request, _('شما اجازه دسترسی به این محصول را ندارید.'))
        return redirect('products:product_list')
    
    # محاسبه آمار فروش محصول
    from invoices.models import InvoiceItem
    stats = {
        'quantity_sold': InvoiceItem.objects.filter(product=product).aggregate(Sum('quantity'))['quantity__sum'] or 0,
        'total_sales': InvoiceItem.objects.filter(product=product).aggregate(total_revenue=Sum(F('quantity') * F('price')))['total_revenue'] or 0,
        'invoice_count': InvoiceItem.objects.filter(product=product).count(),
    }
    
    # گرفتن آیتم‌های فاکتور مربوط به این محصول
    sales_items = InvoiceItem.objects.filter(product=product).select_related('invoice').order_by('-invoice__created_at')
    
    context = {
        'product': product,
        'stats': stats,
        'sales_data': {
            'items': sales_items
        },
        'title': product.name
    }
    return render(request, 'products/product_detail.html', context)

@login_required
def product_create(request, booth_id=None):
    """ایجاد محصول جدید"""
    # بررسی اینکه آیا یک غرفه خاص برای محصول تعیین شده است
    booth = None
    if booth_id:
        booth = get_object_or_404(Booth, pk=booth_id)
    
    # فقط ادمین‌ها می‌توانند محصول جدید اضافه کنند
    if not request.user.is_staff:
        messages.error(request, _('فقط ادمین می‌تواند محصول جدید اضافه کند.'))
        if booth:
            return redirect('booth_detail', pk=booth.pk)
        else:
            return redirect('products:product_list')
    
    if request.method == 'POST':
        # بررسی اینکه آیا غرفه ارسال شده است
        post_data = request.POST.copy()  # ساخت کپی از داده‌های ارسالی برای تغییر آن
        
        # اگر booth_id وجود دارد، به صورت دستی اضافه شود
        if booth_id and booth and 'booth' not in post_data:
            post_data['booth'] = booth.id
        
        form = ProductForm(post_data, request.FILES, booth=booth)
        if form.is_valid():
            try:
                # اطمینان از تنظیم غرفه قبل از ذخیره
                product = form.save(commit=False)
                if booth_id and booth:
                    product.booth = booth
                product.save()
                
                messages.success(request, _('محصول با موفقیت ایجاد شد.'))
                
                # اگر از طریق صفحه غرفه ایجاد شده، به صفحه غرفه بازگردید
                if booth:
                    return redirect('booth_detail', pk=booth.pk)
                return redirect('products:product_detail', pk=product.pk)
            except Exception as e:
                messages.error(request, _('خطا در ذخیره محصول: {}'.format(str(e))))
        else:
            # نمایش خطاهای فرم به صورت مناسب
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, _('خطا در فیلد {}: {}'.format(form.fields[field].label, error)))
    else:
        form = ProductForm(booth=booth)
    
    context = {
        'form': form,
        'booth': booth,
        'title': _('ایجاد محصول جدید')
    }
    return render(request, 'products/product_form.html', context)

@login_required
def product_update(request, pk):
    """ویرایش محصول"""
    product = get_object_or_404(Product, pk=pk)
    
    # بررسی دسترسی کاربر به محصول - فقط ادمین
    if not request.user.is_staff:
        messages.error(request, _('فقط ادمین می‌تواند محصول را ویرایش کند.'))
        return redirect('products:product_detail', pk=product.pk)
    
    if request.method == 'POST':
        form = ProductForm(request.POST, request.FILES, instance=product, booth=product.booth)
        if form.is_valid():
            form.save()
            messages.success(request, _('محصول با موفقیت بروزرسانی شد.'))
            return redirect('products:product_detail', pk=product.pk)
    else:
        form = ProductForm(instance=product, booth=product.booth)
    
    context = {
        'form': form,
        'product': product,
        'title': _('ویرایش محصول')
    }
    return render(request, 'products/product_form.html', context)

@login_required
def product_delete(request, pk):
    """حذف محصول"""
    product = get_object_or_404(Product, pk=pk)
    
    # بررسی دسترسی کاربر به محصول - فقط ادمین
    if not request.user.is_staff:
        messages.error(request, _('فقط ادمین می‌تواند محصول را حذف کند.'))
        return redirect('products:product_detail', pk=product.pk)
    
    if request.method == 'POST':
        booth = product.booth  # ذخیره غرفه قبل از حذف
        product.delete()
        messages.success(request, _('محصول با موفقیت حذف شد.'))
        return redirect('booth_detail', pk=booth.pk)
    
    context = {
        'product': product,
        'title': _('حذف محصول')
    }
    return render(request, 'products/product_confirm_delete.html', context)

@login_required
def product_bulk_upload(request):
    """آپلود گروهی محصولات از طریق فایل CSV"""
    # بررسی دسترسی کاربر
    if not request.user.is_staff and not Booth.objects.filter(manager=request.user).exists():
        messages.error(request, _('شما اجازه آپلود گروهی محصولات را ندارید.'))
        return redirect('products:product_list')
    
    if request.method == 'POST':
        form = ProductBulkUploadForm(request.POST, request.FILES)
        if form.is_valid():
            csv_file = request.FILES['file']
            booth = form.cleaned_data['booth']
            
            # بررسی دسترسی کاربر به غرفه
            if not request.user.is_staff and request.user != booth.manager and request.user not in booth.staff.all():
                messages.error(request, _('شما اجازه آپلود محصولات برای این غرفه را ندارید.'))
                return redirect('products:product_list')
            
            # پردازش فایل CSV
            try:
                decoded_file = csv_file.read().decode('utf-8')
                io_string = io.StringIO(decoded_file)
                reader = csv.DictReader(io_string)
                
                products_created = 0
                products_failed = 0
                
                for row in reader:
                    try:
                        # بررسی فیلدهای اجباری
                        if not row.get('name') or not row.get('code') or not row.get('price'):
                            products_failed += 1
                            continue
                        
                        # یافتن یا ایجاد واحد
                        unit = None
                        if row.get('unit'):
                            unit, created = UnitOfMeasure.objects.get_or_create(
                                name=row['unit'],
                                defaults={'symbol': row.get('unit_symbol', row['unit'][:2])}
                            )
                        
                        # ایجاد محصول
                        Product.objects.create(
                            code=row['code'],
                            name=row['name'],
                            description=row.get('description', ''),
                            booth=booth,
                            price=float(row['price']),
                            unit=unit,
                            is_available=True
                        )
                        products_created += 1
                    except Exception:
                        products_failed += 1
                
                messages.success(
                    request,
                    _('آپلود محصولات انجام شد. {0} محصول ایجاد شده، {1} محصول ناموفق.').format(
                        products_created, products_failed
                    )
                )
                return redirect('booth_detail', pk=booth.pk)
            except Exception as e:
                messages.error(request, _('خطا در پردازش فایل CSV: {0}').format(str(e)))
    else:
        form = ProductBulkUploadForm()
    
    # دریافت لیست غرفه‌ها با توجه به دسترسی کاربر
    if request.user.is_staff:
        # ادمین‌ها به همه غرفه‌ها دسترسی دارند
        available_booths = Booth.objects.all().order_by('name')
    else:
        # کاربران عادی فقط به غرفه‌هایی که مدیر یا کارمند آن‌ها هستند دسترسی دارند
        managed_booths = Booth.objects.filter(manager=request.user)
        staff_booths = Booth.objects.filter(staff=request.user)
        available_booths = (managed_booths | staff_booths).distinct().order_by('name')
    
    # دریافت لیست واحدهای اندازه‌گیری
    available_units = UnitOfMeasure.objects.all().order_by('name')
    
    context = {
        'form': form,
        'title': _('افزودن گروهی محصول'),
        'available_booths': available_booths,
        'available_units': available_units
    }
    return render(request, 'products/product_bulk_upload.html', context)

@login_required
def unit_list(request):
    """نمایش لیست واحدهای اندازه‌گیری"""
    units = UnitOfMeasure.objects.all()
    
    context = {
        'units': units,
        'title': _('واحدهای اندازه‌گیری')
    }
    return render(request, 'products/unit_list.html', context)

@login_required
def unit_create(request):
    """ایجاد واحد اندازه‌گیری جدید"""
    if not request.user.is_staff:
        messages.error(request, _('شما اجازه ایجاد واحد اندازه‌گیری ندارید.'))
        return redirect('products:unit_list')
    
    if request.method == 'POST':
        form = UnitOfMeasureForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, _('واحد اندازه‌گیری با موفقیت ایجاد شد.'))
            return redirect('products:unit_list')
    else:
        form = UnitOfMeasureForm()
    
    context = {
        'form': form,
        'title': _('ایجاد واحد اندازه‌گیری جدید')
    }
    return render(request, 'products/unit_form.html', context)

@login_required
def unit_update(request, pk):
    """ویرایش واحد اندازه‌گیری"""
    unit = get_object_or_404(UnitOfMeasure, pk=pk)
    
    if not request.user.is_staff:
        messages.error(request, _('شما اجازه ویرایش واحد اندازه‌گیری را ندارید.'))
        return redirect('products:unit_list')
    
    if request.method == 'POST':
        form = UnitOfMeasureForm(request.POST, instance=unit)
        if form.is_valid():
            form.save()
            messages.success(request, _('واحد اندازه‌گیری با موفقیت بروزرسانی شد.'))
            return redirect('products:unit_list')
    else:
        form = UnitOfMeasureForm(instance=unit)
    
    context = {
        'form': form,
        'unit': unit,
        'title': _('ویرایش واحد اندازه‌گیری')
    }
    return render(request, 'products/unit_form.html', context)

@login_required
def unit_delete(request, pk):
    """حذف واحد اندازه‌گیری"""
    unit = get_object_or_404(UnitOfMeasure, pk=pk)
    
    if not request.user.is_staff:
        messages.error(request, _('شما اجازه حذف واحد اندازه‌گیری را ندارید.'))
        return redirect('products:unit_list')
    
    # بررسی اینکه آیا محصولی از این واحد استفاده می‌کند
    if Product.objects.filter(unit=unit).exists():
        messages.error(request, _('این واحد اندازه‌گیری قابل حذف نیست زیرا در محصولات استفاده شده است.'))
        return redirect('products:unit_list')
    
    if request.method == 'POST':
        unit.delete()
        messages.success(request, _('واحد اندازه‌گیری با موفقیت حذف شد.'))
        return redirect('products:unit_list')
    
    context = {
        'unit': unit,
        'title': _('حذف واحد اندازه‌گیری')
    }
    return render(request, 'products/unit_confirm_delete.html', context)

@login_required
def create_unit_ajax(request):
    """ایجاد واحد اندازه‌گیری جدید با Ajax"""
    if request.method == 'POST':
        form = UnitOfMeasureForm(request.POST)
        if form.is_valid():
            unit = form.save()
            return JsonResponse({
                'success': True,
                'id': unit.id,
                'name': unit.name,
                'message': _('واحد اندازه‌گیری با موفقیت ایجاد شد')
            })
        else:
            errors = {}
            for field, error_list in form.errors.items():
                errors[field] = [str(error) for error in error_list]
            return JsonResponse({
                'success': False,
                'errors': errors
            })
    return JsonResponse({'success': False, 'message': _('درخواست نامعتبر')})

@login_required
def get_units_json(request):
    """دریافت لیست واحدهای اندازه‌گیری به صورت JSON برای Select2"""
    query = request.GET.get('q', '')
    units = UnitOfMeasure.objects.all()
    
    if query:
        units = units.filter(name__icontains=query)
    
    results = []
    for unit in units:
        results.append({
            'id': unit.id,
            'text': f"{unit.name} ({unit.symbol})"
        })
    
    return JsonResponse({'results': results})

@login_required
def product_toggle_availability(request, pk):
    """تغییر وضعیت موجود/ناموجود بودن محصول"""
    product = get_object_or_404(Product, pk=pk)
    
    # بررسی دسترسی کاربر به محصول
    if not request.user.is_staff and request.user != product.booth.manager and request.user not in product.booth.staff.all():
        messages.error(request, _('شما اجازه تغییر وضعیت این محصول را ندارید.'))
        return redirect('products:product_detail', pk=product.pk)
    
    if request.method == 'POST':
        is_available = request.POST.get('is_available', '') == 'true'
        product.is_available = is_available
        product.save()
        
        status_message = _('موجود') if is_available else _('ناموجود')
        messages.success(request, _('وضعیت محصول به {} تغییر یافت.').format(status_message))
    
    return redirect('products:product_detail', pk=product.pk)

@login_required
def product_api(request, pk):
    """API برای دریافت اطلاعات یک محصول"""
    product = get_object_or_404(Product, pk=pk)
    
    # اطمینان از دسترسی کاربر به داده‌های محصول
    if not request.user.is_staff and request.user != product.booth.manager and request.user not in product.booth.staff.all():
        return JsonResponse({'error': 'دسترسی غیر مجاز'}, status=403)
    
    # تبدیل به فرمت JSON
    product_data = {
        'id': product.id,
        'name': product.name,
        'code': product.code,
        'price': product.price,
        'is_available': product.is_available,
        'unit': product.unit.name if product.unit else '',
        'booth_id': product.booth.id,
        'booth_name': product.booth.name,
        'image': product.get_image_url() if hasattr(product, 'get_image_url') else None,
        'description': product.description if hasattr(product, 'description') else '',
        'discount_percent': product.discount_percent if hasattr(product, 'discount_percent') else 0,
    }
    
    return JsonResponse(product_data, safe=False)

@login_required
def product_create_api(request):
    """ثبت محصول جدید از طریق API"""
    if request.method != 'POST':
        return JsonResponse({'success': False, 'error': 'درخواست اشتباه'}, status=400)
    
    try:
        # دریافت داده‌های محصول از درخواست
        data = request.POST
        
        # بررسی قیمت و تبدیل به عدد
        try:
            price = data.get('price', '0')
            if price:
                price = price.replace(',', '')  # حذف کاماها
            price = float(price)
        except (ValueError, TypeError):
            return JsonResponse({'success': False, 'error': 'قیمت وارد شده معتبر نیست'}, status=400)
        
        # بررسی موجود بودن غرفه
        try:
            booth_id = data.get('booth', '')
            if booth_id:
                booth = Booth.objects.get(pk=booth_id)
            else:
                return JsonResponse({'success': False, 'error': 'غرفه انتخاب نشده است'}, status=400)
        except Booth.DoesNotExist:
            return JsonResponse({'success': False, 'error': 'غرفه انتخاب شده وجود ندارد'}, status=400)
        
        # بررسی موجود بودن واحد
        try:
            unit_id = data.get('unit', '')
            if unit_id:
                unit = UnitOfMeasure.objects.get(pk=unit_id)
            else:
                return JsonResponse({'success': False, 'error': 'واحد انتخاب نشده است'}, status=400)
        except UnitOfMeasure.DoesNotExist:
            return JsonResponse({'success': False, 'error': 'واحد انتخاب شده وجود ندارد'}, status=400)
        
        # بررسی کد محصول تکراری
        code = data.get('code', '')
        if code and Product.objects.filter(code=code).exists():
            return JsonResponse({'success': False, 'error': 'کد محصول تکراری است'}, status=400)
        
        # بررسی نام محصول
        name = data.get('name', '')
        if not name:
            return JsonResponse({'success': False, 'error': 'نام محصول الزامی است'}, status=400)
            
        # ایجاد محصول جدید
        product = Product.objects.create(
            code=code,
            name=name,
            price=price,
            unit=unit,
            booth=booth,
            description=data.get('description', ''),
            is_available=True  # طبق درخواست، موجودی همیشه فعال باشد
        )
        
        # پردازش تصویر در صورت وجود
        if 'image_data' in data and data['image_data']:
            try:
                import base64
                from django.core.files.base import ContentFile
                
                # تبدیل Base64 به فایل
                image_data = data['image_data']
                image_name = data.get('image_name', 'product.jpg')
                
                # ایجاد فایل از داده Base64
                image_content = ContentFile(base64.b64decode(image_data))
                
                # ذخیره در فیلد تصویر محصول
                product.image.save(image_name, image_content, save=True)
                
                # پردازش تصویر توسط متدهای داخلی مدل رخ می‌دهد (process_image)
                # در متد save مدل Product
            except Exception as e:
                # در صورت خطا در پردازش تصویر، محصول بدون تصویر ذخیره می‌شود
                return JsonResponse({
                    'success': True, 
                    'product_id': product.id,
                    'message': 'محصول با موفقیت ثبت شد، اما در پردازش تصویر خطا رخ داد: ' + str(e)
                })
        
        # بازگرداندن پاسخ موفق
        return JsonResponse({
            'success': True, 
            'product_id': product.id,
            'message': 'محصول با موفقیت ثبت شد'
        })
        
    except Exception as e:
        # در صورت بروز خطا
        return JsonResponse({
            'success': False, 
            'error': f'خطای سیستمی: {str(e)}'
        }, status=500)

@login_required
def product_excel_template(request):
    """تولید فایل اکسل نمونه برای آپلود دسته‌ای محصولات
    این فایل شامل سه شیت است:
    1. شیت نمونه داده‌ها (ساختار اصلی داده‌ها)
    2. شیت لیست غرفه‌ها
    3. شیت لیست واحدهای اندازه‌گیری
    """
    # دریافت لیست غرفه‌ها و واحدهای اندازه‌گیری از پایگاه داده
    from booths.models import Booth
    
    # برای غیر ادمین، فقط غرفه‌هایی که مدیر یا کارمند آنهاست را نشان دهید
    if request.user.is_staff:
        booths = Booth.objects.all().order_by('name')
    else:
        managed_booths = Booth.objects.filter(manager=request.user)
        staff_booths = Booth.objects.filter(staff=request.user)
        booths = (managed_booths | staff_booths).distinct().order_by('name')
    
    units = UnitOfMeasure.objects.all().order_by('name')
    
    # ایجاد فایل اکسل جدید
    workbook = openpyxl.Workbook()
    
    # تنظیم استایل‌ها
    header_font = Font(name='B Nazanin', bold=True, size=12)
    header_fill = PatternFill(start_color='E6E6E6', end_color='E6E6E6', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    header_border = Border(
        left=Side(style='thin'), 
        right=Side(style='thin'), 
        top=Side(style='thin'), 
        bottom=Side(style='thin')
    )
    
    data_font = Font(name='B Nazanin', size=11)
    data_alignment = Alignment(horizontal='center', vertical='center')
    data_border = Border(
        left=Side(style='thin'), 
        right=Side(style='thin'), 
        top=Side(style='thin'), 
        bottom=Side(style='thin')
    )
    
    # 1. شیت نمونه داده‌ها
    main_sheet = workbook.active
    main_sheet.title = "نمونه محصولات"
    
    # ستون‌های اصلی
    headers = [
        "کد کالا", 
        "نام کالا", 
        "قیمت", 
        "واحد", 
        "توضیحات", 
        "غرفه"
    ]
    
    # افزودن هدرها
    for col_idx, header in enumerate(headers, 1):
        cell = main_sheet.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = header_border
        main_sheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15
    
    # افزودن چند ردیف نمونه
    sample_data = [
        ["P001", "شامپو شینیون", "120000", "عدد", "شامپوی مناسب برای موهای چرب", "لوازم بهداشتی"],
        ["P002", "کتاب آشپزی", "85000", "عدد", "شامل دستور پخت غذاهای ایرانی", "کتاب"],
        ["P003", "زعفران", "1500000", "گرم", "زعفران خراسان درجه یک", "مواد غذایی"],
    ]
    
    for row_idx, row_data in enumerate(sample_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = main_sheet.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = data_border
    
    # 2. شیت لیست غرفه‌ها
    booth_sheet = workbook.create_sheet(title="غرفه‌ها")
    
    # هدر غرفه‌ها
    cell = booth_sheet.cell(row=1, column=1)
    cell.value = "نام غرفه"
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = header_border
    booth_sheet.column_dimensions['A'].width = 30
    
    # افزودن اسامی غرفه‌ها
    for idx, booth in enumerate(booths, 2):
        cell = booth_sheet.cell(row=idx, column=1)
        cell.value = booth.name
        cell.font = data_font
        cell.alignment = data_alignment
        cell.border = data_border
    
    # 3. شیت لیست واحدها
    unit_sheet = workbook.create_sheet(title="واحدها")
    
    # هدر واحدها
    cell = unit_sheet.cell(row=1, column=1)
    cell.value = "نام واحد"
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = header_border
    unit_sheet.column_dimensions['A'].width = 20
    
    # افزودن اسامی واحدها
    for idx, unit in enumerate(units, 2):
        cell = unit_sheet.cell(row=idx, column=1)
        cell.value = unit.name
        cell.font = data_font
        cell.alignment = data_alignment
        cell.border = data_border
    
    # ذخیره فایل و ارسال پاسخ
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=product_template.xlsx'
    
    # ذخیره فایل در پاسخ HTTP
    workbook.save(response)
    return response
